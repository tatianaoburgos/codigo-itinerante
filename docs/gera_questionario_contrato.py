"""Gera o questionario de intake do contrato de permuta (.docx).

Lista de perguntas que a autora responde antes de redigir cada contrato real
(uma pergunta por dado que MUDA de cliente para cliente). As clausulas fixas
(mesmas em todo contrato) ficam so de referencia no topo, sem pergunta.

O valor de R$3.000 usado para calcular as noites e parametro interno da
autora e nao deve aparecer no texto do contrato entregue ao cliente -- so o
numero de noites resultante aparece.

Uso:
    pip install python-docx   # se ainda nao instalado
    python gera_questionario_contrato.py

Saida: questionario-contrato.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AQUI = Path(__file__).parent

TINTA = RGBColor(0x1A, 0x19, 0x15)
QUEIMADO = RGBColor(0xB8, 0x5C, 0x2E)
CINZA = RGBColor(0x6B, 0x66, 0x5E)

CLAUSULAS_FIXAS = [
    "Piso de seguranca: nunca menos de 20 noites, mesmo com diaria alta.",
    "4 rodadas de ajuste durante a construcao, antes da entrega.",
    "15 dias pos-entrega: 1 manutencao/mudanca pequena gratis "
    "(texto/foto pontual). Mudanca grande (secao nova, layout) e sempre "
    "cobranca a parte, mesmo dentro dos 15 dias.",
    "Apos os 15 dias: qualquer mudanca e novo acordo (contrata a autora ou "
    "outra pessoa).",
    "Dominio: registrado pelo proprio cliente (CPF/CNPJ dele), titular "
    "desde o inicio.",
    "Portabilidade do codigo: gratis, uma vez, dentro de 30 dias a partir "
    "da data em que a autora for embora do hostel. Codigo exportado e para "
    "uso no site do cliente -- nao pode ser revendido/reusado como motor "
    "para outros clientes.",
    "Se o material chegar sem 10 dias de antecedencia da chegada da "
    "autora: ela viaja mesmo assim; o trabalho passa a contar a partir de "
    "quando o material chegar.",
    "Esgotadas as rodadas de ajuste, o site fica como entregue e as "
    "noites continuam devidas.",
]

GRUPOS: list[tuple[str, list[str]]] = [
    (
        "Identificacao",
        [
            "Nome do hostel (como aparece no site)",
            "Nome completo do responsavel legal (quem assina)",
            "CPF ou CNPJ do responsavel",
            "Endereco completo do hostel",
            "E-mail de contato",
            "WhatsApp de contato",
        ],
    ),
    (
        "Calculo da contrapartida (uso interno -- nao aparece no contrato)",
        [
            "Diaria real do hostel usada no calculo (R$)",
            "Numero de noites resultante (3.000 / diaria, minimo 20)",
            "Numero de noites final combinado, se diferente do calculado "
            "-- e por que",
        ],
    ),
    (
        "Acomodacao durante a estadia",
        [
            "Tipo de quarto/cama oferecido (privativo, dormitorio misto, "
            "dormitorio feminino, outro)",
            "Quarto/cama especifico garantido (fixo, sem remanejamento)",
        ],
    ),
    (
        "Janela de uso e datas",
        [
            "Periodo/temporada em que as noites podem ser usadas (ex.: "
            "baixa temporada, meses especificos)",
            "Antecedencia minima exigida para reservar",
            "Prazo de validade das noites, se houver",
            "Data prevista de chegada da autora no hostel",
            "Data prevista de saida",
        ],
    ),
    (
        "Prazos do site",
        [
            "Data limite para o cliente enviar fotos/material "
            "(10 dias antes da chegada)",
            "Data prevista de entrega final do site",
        ],
    ),
    (
        "Dominio",
        [
            "Dominio pretendido (ex.: nomedohostel.com.br), se ja souber",
        ],
    ),
    (
        "Formalizacao",
        [
            "Assinatura eletronica (gov.br) ou papel fisico -- qual dos "
            "dois",
            "Nome completo e e-mail para a assinatura (se eletronica)",
        ],
    ),
    (
        "Observacoes",
        [
            "Condicoes especiais negociadas so nesse contrato, fora do "
            "padrao",
        ],
    ),
]


def configurar_estilos(documento: Document) -> None:
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TINTA
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3


def titulo_principal(documento: Document) -> None:
    p = documento.add_paragraph()
    r = p.add_run("Questionário — intake do contrato de permuta")
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = TINTA

    sub = documento.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    r_sub = sub.add_run(
        "Preencher antes de redigir cada contrato novo. Cobre só o que "
        "muda de cliente para cliente — as cláusulas abaixo já são padrão "
        "e não precisam de resposta."
    )
    r_sub.font.italic = True
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = CINZA


def secao_clausulas_fixas(documento: Document) -> None:
    h = documento.add_paragraph()
    h.paragraph_format.space_before = Pt(4)
    r = h.add_run("Cláusulas padrão (referência, sem pergunta)")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = QUEIMADO

    for item in CLAUSULAS_FIXAS:
        p = documento.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.size = Pt(10)
        run.font.color.rgb = CINZA

    espaco = documento.add_paragraph()
    espaco.paragraph_format.space_after = Pt(10)


def secao_perguntas(documento: Document, grupos: list[tuple[str, list[str]]]) -> None:
    contador = 1
    for titulo_grupo, perguntas in grupos:
        h = documento.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(titulo_grupo)
        r.font.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = QUEIMADO

        for pergunta in perguntas:
            p = documento.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r_num = p.add_run(f"{contador}. ")
            r_num.font.bold = True
            r_num.font.color.rgb = TINTA
            r_txt = p.add_run(pergunta)
            r_txt.font.color.rgb = TINTA

            resposta = documento.add_paragraph()
            resposta.paragraph_format.space_after = Pt(8)
            resposta.paragraph_format.left_indent = Pt(18)
            r_resp = resposta.add_run("R: " + "_" * 60)
            r_resp.font.color.rgb = CINZA
            r_resp.font.size = Pt(10)

            contador += 1


def main() -> None:
    documento = Document()
    configurar_estilos(documento)
    titulo_principal(documento)
    secao_clausulas_fixas(documento)
    secao_perguntas(documento, GRUPOS)

    destino = AQUI / "questionario-contrato.docx"
    documento.save(str(destino))
    print(f"gerado: {destino}")


if __name__ == "__main__":
    main()
