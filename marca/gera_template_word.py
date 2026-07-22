"""Gera o template Word (.docx) com a assinatura da marca Código Itinerante.

Sistema visual conforme docs/miv.md secao 5 ("Documentos"): cabecalho com
wordmark + tipo do documento, rodape com monograma + contato + paginacao,
fundo Papel, corpo em Archivo com destaques em Queimado. E o arquivo-base
generico que a autora duplica para propostas/briefings/contratos reais —
nao inclui texto comercial de verdade, so conteudo placeholder.

Uso:
    python prepara_fontes_documento.py   # uma vez, instala as fontes
    pip install python-docx svglib reportlab rlPyCairo pywin32
    python gera_template_word.py

Saida: template-documento.docx
"""

import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from reportlab.graphics import renderPM
from svglib.svglib import svg2rlg

AQUI = Path(__file__).parent

PAPEL = 0xFBF9F4
TINTA = RGBColor(0x1A, 0x19, 0x15)
QUEIMADO = RGBColor(0xB8, 0x5C, 0x2E)
CINZA_PEDRA = RGBColor(0x97, 0x92, 0x8A)

FONTE_CORPO = "Archivo Doc"
FONTE_MONO = "Martian Mono Doc"


def rasterizar_svg(caminho_svg: Path, destino_png: Path, dpi: int = 300) -> None:
    """Converte um SVG em PNG com fundo Papel, para embutir no docx."""
    desenho = svg2rlg(str(caminho_svg))
    renderPM.drawToFile(desenho, str(destino_png), fmt="PNG", bg=PAPEL, dpi=dpi)


def borda_paragrafo(paragrafo, lado: str, cor_hex: str, tamanho_oitavos_pt: int, espaco_pt: int) -> None:
    """Adiciona uma regua (borda simples) a um lado do paragrafo via OOXML."""
    pPr = paragrafo._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    borda = OxmlElement(f"w:{lado}")
    borda.set(qn("w:val"), "single")
    borda.set(qn("w:sz"), str(tamanho_oitavos_pt))
    borda.set(qn("w:space"), str(espaco_pt))
    borda.set(qn("w:color"), cor_hex)
    pBdr.append(borda)
    pPr.append(pBdr)


def campo_simples(paragrafo, instrucao: str, texto_cache: str) -> None:
    """Insere um campo de Word (ex.: PAGE, NUMPAGES) com um valor de cache exibido."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instrucao)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = texto_cache
    r.append(t)
    fld.append(r)
    paragrafo._p.append(fld)


def fundo_pagina(documento: Document, cor_hex: str) -> None:
    """Define a cor de fundo da pagina (elemento w:background, filho de w:document)."""
    fundo = OxmlElement("w:background")
    fundo.set(qn("w:color"), cor_hex)
    documento.element.insert(0, fundo)


def configurar_secao(documento: Document) -> Emu:
    """A4, margens de 2,5cm (2cm em cima/baixo). Retorna a largura util em EMU."""
    secao = documento.sections[0]
    secao.page_width = Cm(21)
    secao.page_height = Cm(29.7)
    secao.left_margin = secao.right_margin = Cm(2.5)
    secao.top_margin = secao.bottom_margin = Cm(2)
    return Emu(secao.page_width - secao.left_margin - secao.right_margin)


def configurar_estilos(documento: Document) -> None:
    normal = documento.styles["Normal"]
    normal.font.name = FONTE_CORPO
    normal.font.size = Pt(11)
    normal.font.color.rgb = TINTA
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.3


def construir_cabecalho(documento: Document, largura_util: Emu, wordmark_png: Path, rotulo: str) -> None:
    secao = documento.sections[0]
    p = secao.header.paragraphs[0]
    p.style = documento.styles["Normal"]
    p.paragraph_format.tab_stops.add_tab_stop(largura_util, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run().add_picture(str(wordmark_png), height=Cm(0.9))
    r = p.add_run()
    r.add_tab()
    r.add_text(rotulo)
    r.font.name = FONTE_MONO
    r.font.size = Pt(10)
    r.font.color.rgb = QUEIMADO
    borda_paragrafo(p, "bottom", "B85C2E", 16, 6)


def construir_rodape(documento: Document, largura_util: Emu, monograma_png: Path, site: str) -> None:
    secao = documento.sections[0]
    p = secao.footer.paragraphs[0]
    p.style = documento.styles["Normal"]
    p.paragraph_format.tab_stops.add_tab_stop(Emu(largura_util // 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(largura_util, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run().add_picture(str(monograma_png), height=Cm(0.6))

    r_site = p.add_run()
    r_site.add_tab()
    r_site.add_text(site)
    r_site.font.name = FONTE_MONO
    r_site.font.size = Pt(8)
    r_site.font.color.rgb = CINZA_PEDRA

    r_pag = p.add_run()
    r_pag.add_tab()
    r_pag.font.name = FONTE_MONO
    r_pag.font.size = Pt(8)
    r_pag.font.color.rgb = CINZA_PEDRA
    campo_simples(p, "PAGE", "1")

    r_barra = p.add_run(" / ")
    r_barra.font.name = FONTE_MONO
    r_barra.font.size = Pt(8)
    r_barra.font.color.rgb = CINZA_PEDRA
    campo_simples(p, "NUMPAGES", "1")

    borda_paragrafo(p, "top", "97928A", 4, 6)


def construir_corpo(documento: Document) -> None:
    for rotulo, valor in [
        ("PARA", "[Nome do cliente]"),
        ("DATA", "[00/00/0000]"),
        ("VALIDADE", "[00/00/0000]"),
    ]:
        p = documento.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r_rotulo = p.add_run(rotulo + "  ")
        r_rotulo.font.name = FONTE_MONO
        r_rotulo.font.size = Pt(9)
        r_rotulo.font.color.rgb = CINZA_PEDRA
        r_valor = p.add_run(valor)
        r_valor.font.name = FONTE_CORPO
        r_valor.font.size = Pt(11)
        r_valor.font.color.rgb = TINTA
        r_valor.font.bold = True

    espaco = documento.add_paragraph()
    espaco.paragraph_format.space_after = Pt(6)

    titulo = documento.add_paragraph()
    titulo.paragraph_format.space_before = Pt(6)
    titulo.paragraph_format.space_after = Pt(12)
    r_titulo = titulo.add_run("[Título do documento]")
    r_titulo.font.name = FONTE_CORPO
    r_titulo.font.size = Pt(20)
    r_titulo.font.bold = True
    r_titulo.font.color.rgb = TINTA

    documento.add_paragraph(
        "[Texto do documento. Corpo em Archivo Regular, 11pt, entrelinha 1,3 "
        "— o padrão de leitura confortável em tela e impresso.]"
    )

    destaque = documento.add_paragraph()
    destaque.add_run("Um trecho de destaque, como ")
    r_destaque = destaque.add_run("uma cláusula importante ou um valor")
    r_destaque.font.color.rgb = QUEIMADO
    r_destaque.font.bold = True
    destaque.add_run(
        ", aparece em Queimado negrito — sempre um trecho, nunca o parágrafo inteiro."
    )


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        wordmark_png = tmp_path / "wordmark.png"
        monograma_png = tmp_path / "monograma.png"
        rasterizar_svg(AQUI / "wordmark-claro.svg", wordmark_png)
        rasterizar_svg(AQUI / "monograma-claro.svg", monograma_png)

        documento = Document()
        largura_util = configurar_secao(documento)
        configurar_estilos(documento)
        fundo_pagina(documento, "FBF9F4")
        construir_cabecalho(documento, largura_util, wordmark_png, "DOCUMENTO")
        construir_rodape(documento, largura_util, monograma_png, "codigoitinerante.com.br")
        construir_corpo(documento)

        destino = AQUI / "template-documento.docx"
        documento.save(str(destino))
        print(f"gerado: {destino}")


if __name__ == "__main__":
    main()
